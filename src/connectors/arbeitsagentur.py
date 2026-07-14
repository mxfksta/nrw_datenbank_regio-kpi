"""Konnektor Bundesagentur für Arbeit (statistik.arbeitsagentur.de).

**Arbeitslose + Arbeitslosenquoten (Gemeinde-/Kreisebene)** → Cluster
"Arbeitslosigkeit (BA)". Quelle ist EINE bundesweite ZIP ("dlk" =
Deutschland/Länder/Kreise) mit zwei XLSX (Bestand + Quoten). Im Blatt
``Übersicht_Kreise`` steht je Zeile ein Regionalschlüssel; die Spalten sind
Monatswerte. Wir laden die Datei EINMAL pro Lauf (Instanz-Cache) und filtern
daraus alle 7 Regionen — kein Download pro Region, keine Konfiguration nötig
(stabile Default-URL, per ENV ``BA_EINZELHEFT_ZIP_URL`` überschreibbar).
Emittiert wird der jeweils neueste Berichtsmonat mit Wert (``aktuell``);
avg_3j entsteht hier nicht, da die Datei nur ~2 Jahre Monatswerte enthält.

Hinweis: SV-Beschäftigte nach Wirtschaftszweigen (WZ A–U) werden NICHT hier,
sondern über die Landesdatenbank NRW (GENESIS 13111-50i) im statistik_nrw-
Konnektor beschafft — stabiler als der interaktive BA-Report.
"""

from __future__ import annotations

import datetime
import io
import logging
import re
import zipfile
from typing import ClassVar

import openpyxl

from src.config import CLUSTER_ARBEITSLOSIGKEIT, Region
from src.connectors.base import Connector, ConnectorError, NotConfiguredError, SourceLayoutError
from src.models import RawObservation
from src.net import http_get
from src.transform import heute, parse_german_number

log = logging.getLogger(__name__)

QUELLE_NAME = "Statistik der Bundesagentur für Arbeit"
QUELLE_URL_EINZELHEFT = (
    "https://statistik.arbeitsagentur.de (Arbeitslose und Arbeitslosenquoten, Gemeindeebene)"
)

#: Dateinamen-Muster der beiden XLSX in der ZIP → Kennzahl + Einheit
_ARBEITSLOSE_RE = re.compile(r"(?i)arbeitslose(?!nquote)")
_QUOTEN_RE = re.compile(r"(?i)quote")


class ArbeitsagenturConnector(Connector):
    name: ClassVar[str] = "arbeitsagentur"
    phase: ClassVar[int] = 1
    clusters: ClassVar[tuple[str, ...]] = (CLUSTER_ARBEITSLOSIGKEIT,)

    def __init__(self, settings):
        super().__init__(settings)
        # RS → {"Arbeitslose": (stichtag, wert), "Arbeitslosenquote": (stichtag, wert)}
        # Einmal pro Lauf gefüllt (die BA-ZIP ist bundesweit, ~23 MB).
        self._einzelheft_cache: dict[str, dict[str, tuple[str, float]]] | None = None

    def fetch_raw(self, region: Region) -> list[RawObservation]:
        if self._einzelheft_cache is None:
            self._einzelheft_cache = self._load_einzelheft_zip()

        eintrag = self._einzelheft_cache.get(region.regionalschluessel)
        if not eintrag:
            log.warning(
                "BA-Einzelheft: Regionalschlüssel nicht in der bundesweiten Datei",
                extra={"region": region.name, "rs": region.regionalschluessel},
            )
            return []

        stand = heute()
        observations: list[RawObservation] = []
        einheiten = {"Arbeitslose": "Anzahl", "Arbeitslosenquote": "%"}
        for kennzahl, (stichtag, wert) in eintrag.items():
            observations.append(
                RawObservation(
                    region=region.name,
                    regionalschluessel=region.regionalschluessel,
                    kpi_cluster=CLUSTER_ARBEITSLOSIGKEIT,
                    kennzahl=kennzahl,
                    jahr_stichtag=stichtag,
                    wert=wert,
                    einheit=einheiten[kennzahl],
                    quelle_name=QUELLE_NAME,
                    quelle_url=QUELLE_URL_EINZELHEFT,
                    stand_datum=stand,
                )
            )
        return observations

    def _load_einzelheft_zip(self) -> dict[str, dict[str, tuple[str, float]]]:
        """Lädt die bundesweite ZIP und parst beide XLSX (Bestand + Quoten)."""
        url = self.settings.ba_einzelheft_zip_url
        if not url:
            raise NotConfiguredError("BA_EINZELHEFT_ZIP_URL ist leer")
        resp = http_get(url, self.settings)
        cache: dict[str, dict[str, tuple[str, float]]] = {}
        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            namen = zf.namelist()
            arbeitslose = next((n for n in namen if _ARBEITSLOSE_RE.search(n) and n.endswith(".xlsx")), None)
            quoten = next((n for n in namen if _QUOTEN_RE.search(n) and n.endswith(".xlsx")), None)
            if not arbeitslose or not quoten:
                raise SourceLayoutError(
                    f"BA-ZIP: erwartete XLSX (Arbeitslose/Quoten) nicht gefunden, enthalten: {namen}"
                )
            for name, kennzahl in ((arbeitslose, "Arbeitslose"), (quoten, "Arbeitslosenquote")):
                for rs, (stichtag, wert) in self._parse_uebersicht_kreise(zf.read(name)).items():
                    cache.setdefault(rs, {})[kennzahl] = (stichtag, wert)
        return cache

    @staticmethod
    def _parse_uebersicht_kreise(xlsx_bytes: bytes) -> dict[str, tuple[str, float]]:
        """Parst Blatt 'Übersicht_Kreise': RS → (Stichtag, neuester Monatswert).

        Layout: eine Kopfzeile trägt in den Datenspalten Datums-Werte (Monate),
        darunter je Region eine Zeile "<RS> <Name>". 'aktuell' = jüngster Monat
        mit gültigem Wert (Platzhalter/0 der noch nicht berichteten Monate werden
        von rechts übersprungen).
        """
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
        try:
            ws = wb["Übersicht_Kreise"] if "Übersicht_Kreise" in wb.sheetnames else wb.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()

        # Kopfzeile finden: Zeile mit den meisten Datums-Zellen ab Spalte 3
        header_idx = -1
        best = 0
        for i, row in enumerate(rows):
            n = sum(1 for c in row[3:] if isinstance(c, datetime.datetime))
            if n > best:
                best, header_idx = n, i
        if header_idx < 0:
            raise SourceLayoutError(
                "BA-Übersicht_Kreise: keine Datums-Kopfzeile gefunden — Layout geändert?"
            )
        datum_spalten = {
            idx: cell for idx, cell in enumerate(rows[header_idx])
            if isinstance(cell, datetime.datetime)
        }

        rs_re = re.compile(r"^(\d{5})\b")
        ergebnis: dict[str, tuple[str, float]] = {}
        for row in rows[header_idx + 1:]:
            if not row or not isinstance(row[0], str):
                continue
            m = rs_re.match(row[0].strip())
            if not m:
                continue
            rs = m.group(1)
            # jüngste Datumsspalte mit gültigem, positivem Wert
            for idx in sorted(datum_spalten, reverse=True):
                if idx >= len(row):
                    continue
                try:
                    wert = parse_german_number(row[idx])
                except (ValueError, TypeError):
                    continue
                if wert <= 0:
                    continue
                stichtag = datum_spalten[idx].strftime("%Y-%m")
                ergebnis[rs] = (stichtag, wert)
                break
        if not ergebnis:
            raise SourceLayoutError(
                "BA-Übersicht_Kreise: keine Regionszeile (RS + Wert) gefunden — Layout geändert?"
            )
        return ergebnis
